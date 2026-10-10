"""Read candidate primer sequences from the "Oligo" sheet of a Nucleotide List workbook."""

from __future__ import annotations

from dataclasses import dataclass

import openpyxl

_VALID_BASES = set("ACGTNacgtn")


class OligoSheetError(ValueError):
    pass


@dataclass(frozen=True)
class OligoRecord:
    oligo_id: str
    name: str
    sequence: str
    row: int

    @property
    def concat_name(self) -> str:
        """Reproduces the sheet's ``Concat.`` column formula: ``=Name&" ("&ID&")"``."""
        return f"{self.name} ({self.oligo_id})"


def load_oligos(xlsx_path: str, sheet_name: str = "Oligo") -> list[OligoRecord]:
    try:
        wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    except PermissionError as exc:
        raise OligoSheetError(
            f"'{xlsx_path}' を開けませんでした。Excel で開いたままになっていないか確認してください。"
        ) from exc

    if sheet_name not in wb.sheetnames:
        raise OligoSheetError(f"シート '{sheet_name}' が '{xlsx_path}' に見つかりません。")
    ws = wb[sheet_name]

    header_row = next(ws.iter_rows(min_row=1, max_row=1))
    header = {cell.value: idx for idx, cell in enumerate(header_row) if cell.value}
    for required in ("Name", "Sequence"):
        if required not in header:
            raise OligoSheetError(
                f"必須列 '{required}' がヘッダー行に見つかりません: {list(header)}"
            )
    # Column A holds the oligo ID (e.g. "O01-01"); the sheet labels this
    # column "Oligo" as a section header rather than a per-column name, so it
    # is addressed positionally rather than by header lookup.
    id_col = 0
    name_col = header["Name"]
    seq_col = header["Sequence"]

    records: list[OligoRecord] = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        oligo_id = row[id_col].value if id_col < len(row) else None
        name = row[name_col].value if name_col < len(row) else None
        sequence = row[seq_col].value if seq_col < len(row) else None
        if not oligo_id or not name or not sequence:
            continue
        sequence = str(sequence).strip()
        if not sequence or not set(sequence) <= _VALID_BASES:
            continue
        records.append(
            OligoRecord(
                oligo_id=str(oligo_id).strip(),
                name=str(name).strip(),
                sequence=sequence,
                row=row_idx,
            )
        )
    return records
